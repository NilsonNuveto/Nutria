import openpyxl, json, pathlib, sys, datetime
path = sys.argv[1]
w = openpyxl.load_workbook(path, data_only=False)
v = openpyxl.load_workbook(path, data_only=True)
out = pathlib.Path('database/imports'); out.mkdir(parents=True, exist_ok=True)
def value(x):
    if isinstance(x,(datetime.date,datetime.time,datetime.datetime)): return x.isoformat()
    if hasattr(x,'text'): return x.text
    return x
audit = {s.title:[{c.coordinate:value(c.value) for c in row if c.value is not None} for row in s] for s in w}
(out/'workbook-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
foods=[]
for row in v['Alimentos'].iter_rows(min_row=2,max_col=11):
    a=[value(c.value) for c in row]
    if not a[1]: continue
    foods.append(dict(zip(['id','name','category','base_quantity','calories','protein','carbs','fat','fiber','sodium','notes'],[row[1].row-1]+a[1:])))
sources=[{'name':c.value,'url':c.hyperlink.target if c.hyperlink else c.value,'cell':c.coordinate} for row in w['Sites'] for c in row if isinstance(c.value,str) and ('http' in c.value or c.hyperlink)]
data={'foods':foods,'sources':sources,'activities':[{'name':v['Dados'].cell(r,7).value,'factor':v['Dados'].cell(r,8).value,'description':v['Dados'].cell(r,9).value} for r in range(2,7)],'meal_types':[v['Dados'].cell(r,1).value for r in range(2,11)]}
(out/'workbook.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'food_count':len(foods),'sources':sources,'Dieta':[x for x in audit['Dieta'] if x],'Receita':audit['Receita'][-5:],'Dados':audit['Dados'][10:]},ensure_ascii=True))
